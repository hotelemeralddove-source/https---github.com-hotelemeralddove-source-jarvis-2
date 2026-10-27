"""Edit the files JARVIS can read — without ever destroying the original.

Every function here writes to a **new** path next to the source and returns that
path. There is no "overwrite" switch, on purpose: a voice command that mangles a
thesis is not recoverable by apology, and the user's files are the one thing the
project promised never to damage.

What is real, offline and free:

    images   resize, crop, rotate, greyscale, blur, sharpen, brightness, contrast,
             add text, thumbnail, convert format            (Pillow)
    audio    trim, convert, normalise, fade, extract        (ffmpeg)
    video    trim, crop, scale, speed, gif, extract audio,  (ffmpeg)
             extract frames, burn subtitles, watermark
    text     replace, append, prepend                       (pure python)
    docx     find & replace, append paragraphs              (python-docx)
    xlsx     set a cell, append a row, new sheet            (openpyxl)
    pptx     find & replace across slides                   (python-pptx)
    pdf      rotate, split, merge, extract a page range     (pypdf)

What is *not* here: rewriting the words inside a PDF (that needs a layout engine and
silently breaks fonts), and anything that needs a model — image/video *generation*
lives in `generate.py`, which talks to a local backend you run yourself.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess

from . import quiet
import time
from pathlib import Path


def _out(src: str, suffix: str = "", ext: str | None = None) -> Path:
    """A sibling file that cannot collide with the original."""
    src_path = Path(os.path.expanduser(str(src)))
    stamp = time.strftime("%Y%m%d_%H%M%S")
    name = f"{src_path.stem}_{suffix or 'edit'}_{stamp}"
    return src_path.with_name(name + (ext or src_path.suffix))


def _ffmpeg() -> str:
    exe = shutil.which("ffmpeg")
    if not exe:
        raise RuntimeError("ffmpeg is not installed — winget install Gyan.FFmpeg "
                           "(free) and reopen the app")
    return exe


def _run(cmd: list[str], timeout: float = 300) -> tuple[bool, str]:
    try:
        out = quiet.run(cmd, capture_output=True, text=True, timeout=timeout,
                             errors="replace")
        if out.returncode != 0:
            tail = (out.stderr or out.stdout or "").strip().splitlines()[-3:]
            return False, " | ".join(tail) or f"exit {out.returncode}"
        return True, ""
    except Exception as exc:
        return False, f"{type(exc).__name__}: {exc}"


# ------------------------------------------------------------------- images
IMAGE_OPS = ("resize", "crop", "rotate", "greyscale", "grayscale", "blur", "sharpen",
             "brightness", "contrast", "text", "thumbnail", "compress", "convert",
             "flip", "border")


def edit_image(path: str, op: str, **kw) -> str:
    """One operation, one new file. `op` decides which keyword arguments matter."""
    try:
        from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
    except Exception:
        return "ERROR: pillow is not installed (pip install pillow)"
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such image: {src}"
    op = (op or "").lower()
    if op not in IMAGE_OPS:
        return f"ERROR: unknown image op '{op}'. I can do: {', '.join(IMAGE_OPS)}"
    try:
        img = Image.open(src)
        img.load()
        if op == "resize":
            w = int(kw.get("width") or 0)
            h = int(kw.get("height") or 0)
            if not (w or h):
                return "ERROR: resize needs width and/or height"
            if not w:
                w = round(img.width * h / img.height)
            if not h:
                h = round(img.height * w / img.width)
            out = img.resize((max(1, w), max(1, h)), Image.LANCZOS)
        elif op == "thumbnail":
            size = int(kw.get("size") or 512)
            out = img.copy()
            out.thumbnail((size, size), Image.LANCZOS)
        elif op == "crop":
            x, y = int(kw.get("x") or 0), int(kw.get("y") or 0)
            w = int(kw.get("width") or img.width - x)
            h = int(kw.get("height") or img.height - y)
            box = (max(0, x), max(0, y), min(img.width, x + w), min(img.height, y + h))
            if box[2] <= box[0] or box[3] <= box[1]:
                return "ERROR: crop area is empty"
            out = img.crop(box)
        elif op == "rotate":
            out = img.rotate(float(kw.get("angle") or 90), expand=True)
        elif op == "flip":
            out = img.transpose(Image.FLIP_LEFT_RIGHT if str(kw.get("axis", "h")) == "h"
                                else Image.FLIP_TOP_BOTTOM)
        elif op in ("greyscale", "grayscale"):
            out = img.convert("L")
        elif op == "blur":
            out = img.filter(ImageFilter.GaussianBlur(float(kw.get("radius") or 2.0)))
        elif op == "sharpen":
            out = img.filter(ImageFilter.UnsharpMask(radius=2, percent=180, threshold=3))
        elif op == "brightness":
            factor = float(kw.get("factor") or 1.2)
            out = ImageEnhance.Brightness(img).enhance(max(0.05, factor))
        elif op == "contrast":
            factor = float(kw.get("factor") or 1.2)
            out = ImageEnhance.Contrast(img).enhance(max(0.05, factor))
        elif op == "border":
            width = int(kw.get("width") or 8)
            colour = str(kw.get("colour") or kw.get("color") or "#7fe3f0")
            out = Image.new("RGB", (img.width + 2 * width, img.height + 2 * width), colour)
            out.paste(img.convert("RGB"), (width, width))
        elif op == "text":
            text = str(kw.get("text") or "").strip()
            if not text:
                return "ERROR: add text needs some text"
            out = img.convert("RGB").copy()
            draw = ImageDraw.Draw(out)
            size = max(12, int(kw.get("size") or max(16, out.height // 22)))
            try:
                font = ImageFont.truetype("DejaVuSans-Bold.ttf", size)
            except Exception:
                font = ImageFont.load_default()
            pos = str(kw.get("position") or "bottom")
            margin = max(8, size // 2)
            box = draw.textbbox((0, 0), text, font=font)
            tw, th = box[2] - box[0], box[3] - box[1]
            x = margin
            y = {"top": margin, "center": (out.height - th) // 2,
                 "bottom": out.height - th - margin}.get(pos, out.height - th - margin)
            # a plate behind the text keeps it readable on any background
            draw.rectangle([x - 4, y - 4, x + tw + 4, y + th + 4],
                           fill=str(kw.get("plate") or "#000000"))
            draw.text((x, y), text, font=font, fill=str(kw.get("colour") or "#ffffff"))
        elif op == "compress":
            target = int(kw.get("size") or 1280)
            out = img.convert("RGB")
            if max(out.size) > target:
                out.thumbnail((target, target), Image.LANCZOS)
        elif op == "convert":
            out = img
        else:
            return f"ERROR: '{op}' is not implemented"

        ext = str(kw.get("as") or "").lower() or src.suffix.lower()
        suffix = op
        target = _out(str(src), suffix)
        if op == "convert" or kw.get("as"):
            fmt = {"jpg": ".jpg", "jpeg": ".jpg", "png": ".png", "webp": ".webp",
                   "bmp": ".bmp", "gif": ".gif", "tiff": ".tiff"}.get(ext.lstrip("."), ext)
            if not fmt.startswith("."):
                fmt = "." + fmt
            target = target.with_suffix(fmt)
        if target.suffix.lower() in (".jpg", ".jpeg") and out.mode in ("RGBA", "P", "L"):
            out = out.convert("RGB")
        quality = int(kw.get("quality") or 88)
        if target.suffix.lower() == ".jpg":
            out.save(target, quality=quality, optimize=True)
        else:
            out.save(target)
        return (f"{op} done on {src.name} → {target.name} "
                f"({target.stat().st_size / 1024:.0f} KB, "
                f"{out.width}×{out.height}). The original is untouched.")
    except Exception as exc:
        return f"ERROR: image {op} failed ({type(exc).__name__}: {exc})"


# -------------------------------------------------------------------- audio
AUDIO_OPS = ("trim", "convert", "normalise", "normalize", "fade", "extract", "mono",
             "speed", "volume")


def edit_audio(path: str, op: str, **kw) -> str:
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such audio file: {src}"
    try:
        exe = _ffmpeg()
    except RuntimeError as exc:
        return f"ERROR: {exc}"
    op = (op or "").lower()
    if op not in AUDIO_OPS:
        return f"ERROR: unknown audio op '{op}'. I can do: {', '.join(AUDIO_OPS)}"
    target = _out(str(src), op, ext=str(kw.get("as") or src.suffix))
    cmd = [exe, "-y", "-v", "error", "-i", str(src)]
    if op == "trim":
        start = float(kw.get("start") or 0)
        end = kw.get("end")
        cmd += ["-ss", str(start)]
        if end:
            cmd += ["-to", str(end)]
        cmd += ["-c", "copy"]
    elif op in ("normalise", "normalize"):
        cmd += ["-af", "loudnorm=I=-16:TP=-1.5:LRA=11"]
    elif op == "fade":
        dur = float(kw.get("seconds") or 1.0)
        cmd += ["-af", f"afade=t=in:st=0:d={dur},areverse,afade=t=in:st=0:d={dur},areverse"]
    elif op == "mono":
        cmd += ["-ac", "1"]
    elif op == "speed":
        rate = max(0.25, min(4.0, float(kw.get("factor") or 1.25)))
        cmd += ["-filter:a", f"atempo={rate}"]
    elif op == "volume":
        db = float(kw.get("db") or 3)
        cmd += ["-filter:a", f"volume={db}dB"]
    elif op == "extract":
        pass
    cmd.append(str(target))
    ok, err = _run(cmd)
    if not ok:
        return f"ERROR: ffmpeg failed — {err}"
    return (f"{op} done → {target.name} ({target.stat().st_size / 1024:.0f} KB). "
            f"Original untouched.")


# -------------------------------------------------------------------- video
VIDEO_OPS = ("trim", "crop", "scale", "speed", "gif", "extract_audio", "frames",
             "subtitles", "watermark", "mute", "compress", "rotate", "thumbnail")


def edit_video(path: str, op: str, **kw) -> str:
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such video: {src}"
    try:
        exe = _ffmpeg()
    except RuntimeError as exc:
        return f"ERROR: {exc}"
    op = (op or "").lower()
    if op not in VIDEO_OPS:
        return f"ERROR: unknown video op '{op}'. I can do: {', '.join(VIDEO_OPS)}"
    if op == "frames":
        folder = Path(os.path.expanduser(str(kw.get("into") or src.parent)))
        every = max(1, int(kw.get("every") or 10))
        pattern = folder / f"{src.stem}_frame_%04d.png"
        ok, err = _run([exe, "-y", "-v", "error", "-i", str(src),
                        "-vf", f"fps=1/{every}", str(pattern)])
        if not ok:
            return f"ERROR: frame extraction failed — {err}"
        made = sorted(folder.glob(f"{src.stem}_frame_*.png"))
        return f"extracted {len(made)} frames (one every {every}s) into {folder}"
    if op == "extract_audio":
        target = _out(str(src), "audio", ext=".mp3")
        ok, err = _run([exe, "-y", "-v", "error", "-i", str(src), "-vn",
                        "-acodec", "libmp3lame", "-q:a", "4", str(target)])
        return (f"audio extracted → {target.name}" if ok else f"ERROR: {err}")
    if op == "gif":
        start = float(kw.get("start") or 0)
        dur = float(kw.get("seconds") or 3)
        width = int(kw.get("width") or 480)
        target = _out(str(src), "gif", ext=".gif")
        ok, err = _run([exe, "-y", "-v", "error", "-ss", str(start), "-t", str(dur),
                        "-i", str(src),
                        "-vf", f"fps=12,scale={width}:-1:flags=lanczos,split[a][b];"
                               f"[a]palettegen[p];[b][p]paletteuse", str(target)])
        return (f"GIF made → {target.name}" if ok else f"ERROR: {err}")
    if op == "thumbnail":
        at = float(kw.get("at") or 1)
        target = _out(str(src), "thumb", ext=".png")
        ok, err = _run([exe, "-y", "-v", "error", "-ss", str(at), "-i", str(src),
                        "-frames:v", "1", str(target)])
        return (f"thumbnail → {target.name}" if ok else f"ERROR: {err}")
    if op == "subtitles":
        srt = str(kw.get("srt") or "")
        if not srt or not Path(srt).exists():
            return ("ERROR: subtitles needs an .srt file. I will not invent captions — "
                    "transcribe the audio with a local model first (faster-whisper).")
        target = _out(str(src), "subbed")
        ok, err = _run([exe, "-y", "-v", "error", "-i", str(src), "-i", srt,
                        "-c:v", "libx264", "-crf", "20", "-c:a", "copy",
                        "-vf", f"subtitles={srt}", str(target)])
        return (f"captions burned in → {target.name}" if ok else f"ERROR: {err}")
    if op == "watermark":
        text = str(kw.get("text") or "JARVIS").replace(":", "")
        target = _out(str(src), "marked")
        ok, err = _run([exe, "-y", "-v", "error", "-i", str(src), "-vf",
                        f"drawtext=text='{text}':fontcolor=white@0.8:fontsize=h/22:"
                        f"x=w-tw-20:y=h-th-20", "-c:a", "copy", str(target)])
        return (f"watermark added → {target.name}" if ok else f"ERROR: {err}")
    # filter-based operations produce a new encode
    target = _out(str(src), op)
    cmd = [exe, "-y", "-v", "error", "-i", str(src)]
    if op == "trim":
        start = float(kw.get("start") or 0)
        end = kw.get("end")
        cmd += ["-ss", str(start)]
        if end:
            cmd += ["-to", str(end)]
        cmd += ["-c", "copy"]
    elif op == "scale":
        width = int(kw.get("width") or 1280)
        cmd += ["-vf", f"scale={width}:-2"]
    elif op == "crop":
        w, h = int(kw.get("width") or 0), int(kw.get("height") or 0)
        x, y = int(kw.get("x") or 0), int(kw.get("y") or 0)
        if not (w and h):
            return "ERROR: crop needs width and height"
        cmd += ["-vf", f"crop={w}:{h}:{x}:{y}"]
    elif op == "rotate":
        cmd += ["-vf", f"transpose={1 if int(kw.get('quarter') or 1) > 0 else 2}"]
    elif op == "speed":
        rate = max(0.25, min(4.0, float(kw.get("factor") or 1.5)))
        cmd += ["-filter:v", f"setpts={1 / rate:.4f}*PTS", "-filter:a",
                f"atempo={rate}"]
    elif op == "mute":
        cmd += ["-an", "-c:v", "copy"]
    elif op == "compress":
        crf = max(18, min(40, int(kw.get("crf") or 28)))
        cmd += ["-c:v", "libx264", "-crf", str(crf), "-preset", "veryfast"]
    cmd.append(str(target))
    ok, err = _run(cmd)
    if not ok:
        return f"ERROR: ffmpeg failed — {err}"
    return (f"{op} done → {target.name} "
            f"({target.stat().st_size / 1e6:.1f} MB). Original untouched.")


# ----------------------------------------------------------- documents/text
def edit_text(path: str, op: str, find: str = "", replace: str = "", text: str = "") -> str:
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such file: {src}"
    body = src.read_text(encoding="utf-8", errors="replace")
    if op == "replace":
        if not find:
            return "ERROR: what should I replace?"
        count = body.count(find)
        if not count:
            return f"'{find}' does not appear in {src.name} — nothing changed"
        new = body.replace(find, replace)
    elif op == "append":
        new = body + ("\n" if body and not body.endswith("\n") else "") + text
    elif op == "prepend":
        new = text + "\n" + body
    else:
        return "ERROR: text ops are replace, append, prepend"
    target = _out(str(src), op)
    target.write_text(new, encoding="utf-8")
    return f"{op} done → {target.name}. Original untouched."


def edit_docx(path: str, find: str = "", replace: str = "", append: str = "") -> str:
    try:
        import docx
    except Exception:
        return "ERROR: pip install python-docx"
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such file: {src}"
    try:
        d = docx.Document(str(src))
    except Exception as exc:
        return f"ERROR: could not open {src.name} ({exc}) — .doc is not supported"
    hits = 0
    if find:
        for para in d.paragraphs:
            # runs split words across formatting, so stitch each paragraph first
            if find in para.text:
                hits += para.text.count(find)
                merged = para.text.replace(find, replace)
                for run in list(para.runs)[1:]:
                    run.text = ""
                if para.runs:
                    para.runs[0].text = merged
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if find in cell.text:
                        hits += cell.text.count(find)
                        cell.text = cell.text.replace(find, replace)
    if append:
        for block in append.split("\n\n"):
            d.add_paragraph(block.strip())
    if not hits and not append:
        return f"'{find}' not found in {src.name} — nothing changed"
    target = _out(str(src), "edited")
    d.save(str(target))
    bits = []
    if hits:
        bits.append(f"{hits} replacement(s)")
    if append:
        bits.append(f"{len(append.split(chr(10) * 2))} paragraph(s) added")
    return f"docx edited: {', '.join(bits)} → {target.name}. Original untouched."


def edit_xlsx(path: str, sheet: str = "", cell: str = "", value: str = "",
              row: list | None = None) -> str:
    try:
        import openpyxl
    except Exception:
        return "ERROR: pip install openpyxl"
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such file: {src}"
    try:
        wb = openpyxl.load_workbook(str(src))
    except Exception as exc:
        return f"ERROR: could not open {src.name} ({exc})"
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    changed = ""
    if cell and value != "":
        ws[cell] = _coerce(value)
        changed = f"{ws.title}!{cell} = {value}"
    elif row:
        ws.append([_coerce(str(v)) for v in row])
        changed = f"appended a row to {ws.title}"
    else:
        return "ERROR: give me a cell and value, or a row to append"
    target = _out(str(src), "edited")
    wb.save(str(target))
    return f"spreadsheet edited: {changed} → {target.name}. Original untouched."


def _coerce(value: str):
    text = str(value).strip()
    for caster in (int, float):
        try:
            return caster(text)
        except ValueError:
            continue
    return text


def edit_pptx(path: str, find: str = "", replace: str = "") -> str:
    try:
        from pptx import Presentation
    except Exception:
        return "ERROR: pip install python-pptx"
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such file: {src}"
    try:
        deck = Presentation(str(src))
    except Exception as exc:
        return f"ERROR: could not open {src.name} ({exc})"
    hits = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if find and find in run.text:
                            run.text = run.text.replace(find, replace)
                            hits += 1
    if not hits:
        return f"'{find}' not found in any slide — nothing changed"
    target = _out(str(src), "edited")
    deck.save(str(target))
    return f"slides edited: {hits} run(s) → {target.name}. Original untouched."


def edit_pdf(path: str, op: str = "extract", pages: str = "", angle: int = 90,
             extra: str = "") -> str:
    try:
        from pypdf import PdfReader, PdfWriter
    except Exception:
        return "ERROR: pip install pypdf"
    src = Path(os.path.expanduser(str(path)))
    if not src.exists():
        return f"ERROR: no such file: {src}"
    try:
        reader = PdfReader(str(src))
        writer = PdfWriter()
        if op == "rotate":
            for page in reader.pages:
                page.rotate(int(angle))
                writer.add_page(page)
        elif op == "extract":
            wanted = _page_range(pages, len(reader.pages))
            for i in wanted:
                writer.add_page(reader.pages[i])
        elif op == "merge":
            for page in reader.pages:
                writer.add_page(page)
            for other in [p for p in re.split(r"[;,]", extra) if p.strip()]:
                more = PdfReader(str(Path(os.path.expanduser(other.strip()))))
                for page in more.pages:
                    writer.add_page(page)
        elif op == "split":
            for i, page in enumerate(reader.pages):
                single = PdfWriter()
                single.add_page(page)
                target = _out(str(src), f"page{i + 1}")
                with open(target, "wb") as fh:
                    single.write(fh)
            return f"split into {len(reader.pages)} single-page PDFs next to the original"
        else:
            return "ERROR: pdf ops are extract, rotate, merge, split"
        target = _out(str(src), op)
        with open(target, "wb") as fh:
            writer.write(fh)
        return (f"pdf {op} done → {target.name} "
                f"({len(writer.pages)} page(s)). Original untouched. "
                "Note: text inside a PDF cannot be rewritten safely — export to "
                ".docx if the words need changing.")
    except Exception as exc:
        return f"ERROR: pdf {op} failed ({type(exc).__name__}: {exc})"


def _page_range(spec: str, total: int) -> list[int]:
    """'1-3,7' → [0,1,2,6]. Anything unparseable means 'all pages'."""
    if not spec.strip():
        return list(range(total))
    out: list[int] = []
    for part in re.split(r"[,\s]+", spec.strip()):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            try:
                start, end = int(a), int(b)
            except ValueError:
                continue
            out += [i - 1 for i in range(start, end + 1) if 1 <= i <= total]
        else:
            try:
                i = int(part)
            except ValueError:
                continue
            if 1 <= i <= total:
                out.append(i - 1)
    return sorted(set(out)) or list(range(total))


def describe() -> str:
    return ("I can edit images (resize, crop, rotate, colour, blur, sharpen, text, "
            "thumbnail, convert), audio (trim, convert, normalise, fade, speed, volume), "
            "video (trim, crop, scale, speed, gif, extract audio, frames, subtitles, "
            "watermark, mute, compress, thumbnail), text, docx, xlsx, pptx and pdf.")
