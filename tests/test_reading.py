"""Can ARC actually read every format it claims to read?

The user asked for one thing above all: read anything, then say what it contains,
what is missing and what to improve. Claiming that is easy; these checks build a
real file of every type and assert on what comes back, so the claim is measured.

Two rules encoded here as assertions:

  * a file ARC cannot read must produce a warning, never a made-up summary;
  * a summary must quote the document, never paraphrase it into something new.

Missing optional libraries are reported as SKIP, not FAIL — the honest-degrade
path is part of the design (a user without ffmpeg still gets a working app).

Run:  python tests/test_reading.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import wave
import struct
import math
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from arcdesk import reading as R  # noqa: E402

PASS = 0
FAIL = 0
SKIP = 0


def check(label: str, ok: bool, detail: str = "") -> None:
    global PASS, FAIL
    print(f"   {'OK  ' if ok else 'FAIL'} {label}" + (f"   [{detail}]" if detail else ""))
    if ok:
        PASS += 1
    else:
        FAIL += 1


def skip(label: str, why: str) -> None:
    global SKIP
    print(f"   SKIP {label}   [{why}]")
    SKIP += 1


def build_fixtures(tmp: Path) -> dict[str, Path]:
    """Real files, written by the libraries that read them."""
    made: dict[str, Path] = {}
    txt = tmp / "notes.txt"
    txt.write_text("Revenue grew 18 percent in Q3. Costs rose 6 percent. "
                   "Risks: hiring delay and monsoon logistics.", encoding="utf-8")
    made["txt"] = txt

    csv_file = tmp / "table.csv"
    csv_file.write_text("name,qty,price\nbolt,4,20\nnut,10,5\n", encoding="utf-8")
    made["csv"] = csv_file

    js = tmp / "conf.json"
    js.write_text('{"project": "arc", "tools": [1, 2, 3]}', encoding="utf-8")
    made["json"] = js

    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Quarterly Review", 0)
        doc.add_paragraph("Revenue grew 18% to 4.2 crore rupees in Q3.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Region"
        table.cell(0, 1).text = "Revenue"
        table.cell(1, 0).text = "Karnataka"
        table.cell(1, 1).text = "2.1 crore"
        doc.save(tmp / "report.docx")
        made["docx"] = tmp / "report.docx"
    except Exception:
        pass

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Month", "Units", "Revenue"])
        ws.append(["Jul", 120, 84000])
        ws.append(["Aug", 150, 99000])
        wb.save(tmp / "book.xlsx")
        made["xlsx"] = tmp / "book.xlsx"
    except Exception:
        pass

    try:
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "ARC DESK"
        slide.placeholders[1].text = "Your desktop. Your universe."
        slide.notes_slide.notes_text_frame.text = "say hey jarvis"
        prs.save(tmp / "deck.pptx")
        made["pptx"] = tmp / "deck.pptx"
    except Exception:
        pass

    try:
        from PIL import Image, ImageDraw
        img = Image.new("RGB", (600, 220), (10, 16, 36))
        ImageDraw.Draw(img).text((30, 90), "HELLO ARC DESK 2026", fill=(255, 220, 120))
        img.save(tmp / "shot.png")
        img.resize((200, 80)).save(tmp / "tiny.gif")
        made["png"] = tmp / "shot.png"
        made["gif"] = tmp / "tiny.gif"
    except Exception:
        pass

    try:
        wav = tmp / "tone.wav"
        with wave.open(str(wav), "wb") as handle:
            handle.setnchannels(1)
            handle.setsampwidth(2)
            handle.setframerate(22050)
            handle.writeframes(b"".join(
                struct.pack("<h", int(9000 * math.sin(2 * math.pi * 440 * i / 22050)))
                for i in range(11025)))
        made["wav"] = wav
    except Exception:
        pass

    try:
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with open(tmp / "blank.pdf", "wb") as handle:
            writer.write(handle)
        made["pdf_blank"] = tmp / "blank.pdf"
    except Exception:
        pass

    exe = __import__("shutil").which("ffmpeg")
    if exe:
        clip = tmp / "clip.mp4"
        os.system(f'{exe} -y -v error -f lavfi -i testsrc=size=160x120:rate=10 '
                  f'-t 1 -pix_fmt yuv420p "{clip}" >/dev/null 2>&1')
        if clip.exists():
            made["mp4"] = clip
    return made


def main() -> int:
    print("reading — every format, with honest failures\n")
    tmp = Path(tempfile.mkdtemp(prefix="arcdesk-read-"))
    made = build_fixtures(tmp)
    print(f"   fixtures in {tmp}")

    # 1. detection --------------------------------------------------------
    print("\n1. format detection")
    for key, path in made.items():
        kind = R.kind_of(str(path))
        check(f"{path.name} → {kind}", kind not in ("", "unknown"), kind)

    # 2. text-bearing formats --------------------------------------------
    print("\n2. text extraction")
    doc = R.read(str(made["txt"]))
    check("txt: words extracted", "Revenue" in doc.text and doc.ok)
    check("txt: kind", doc.kind == "text", doc.kind)

    if "docx" in made:
        docx_doc = R.read(str(made["docx"]))
        check("docx: paragraph text", "4.2 crore" in docx_doc.text)
        check("docx: table text included", "Karnataka" in docx_doc.text)
        summary = R.summarise(docx_doc)
        check("docx: summary quotes the file",
              "crore" in summary, summary.splitlines()[0][:48])
    else:
        skip("docx", "python-docx not installed")

    if "xlsx" in made:
        book = R.read(str(made["xlsx"]))
        check("xlsx: cell values", "84000" in book.text and "99000" in book.text)
        stats = R.summarise(book)
        check("xlsx: summary does arithmetic",
              "183,000" in stats or "183000" in stats, "total revenue")
        check("xlsx: names the sheet", "Sales" in stats)
    else:
        skip("xlsx", "openpyxl not installed")

    if "pptx" in made:
        deck = R.read(str(made["pptx"]))
        check("pptx: slide text", "ARC DESK" in deck.text)
        check("pptx: speaker notes", "hey jarvis" in deck.text.lower())
    else:
        skip("pptx", "python-pptx not installed")

    check("csv: rows and header", "bolt" in R.read(str(made["csv"])).text)
    check("json: keys survive", "project" in R.read(str(made["json"])).text)

    # 3. images -----------------------------------------------------------
    print("\n3. images and OCR")
    if "png" in made:
        shot = R.read(str(made["png"]))
        check("png: dimensions", shot.meta.get("width") == 600)
        if shot.text:
            check("png: OCR read the text", "ARC" in shot.text.upper(), shot.text[:40])
        else:
            skip("png OCR", "no tesseract / no text found")
        tiny = R.read(str(made["gif"]))
        check("gif: kind and frames", tiny.kind == "gif" and tiny.meta.get("frames", 1) >= 1)
    else:
        skip("png/gif", "pillow not installed")

    # 4. audio + video ----------------------------------------------------
    print("\n4. audio and video (metadata, not words)")
    if "wav" in made:
        audio = R.read(str(made["wav"]))
        duration = float(audio.meta.get("duration_s") or 0)
        check("wav: duration ≈ 0.5 s", 0.3 <= duration <= 0.8, str(duration))
        check("wav: says it cannot transcribe",
              any("speech-to-text" in w or "metadata only" in w for w in audio.warnings))
        check("wav: does not invent a transcript", audio.text.strip() == "")
    else:
        skip("wav", "stdlib wave unavailable")

    if "mp4" in made:
        video = R.read(str(made["mp4"]))
        check("mp4: dimensions", (video.meta.get("width"), video.meta.get("height")) == (160, 120))
        check("mp4: duration", abs(float(video.meta.get("duration_s") or 0) - 1.0) < 0.3)
    else:
        skip("mp4", "ffmpeg not installed")

    # 5. review -----------------------------------------------------------
    print("\n5. review: missing / improve / add")
    thin = R.read(str(made["txt"]))
    review = R.review(thin)
    check("review has a verdict", "Verdict:" in review)
    check("review names the file", "notes.txt" in review)
    check("review quotes a number from the file", "18" in review or "6" in review)
    if "xlsx" in made:
        book_review = R.review(R.read(str(made["xlsx"])))
        check("xlsx review talks about totals",
              "total" in book_review.lower(), "spreadsheet-specific advice")

    # 6. refusals ---------------------------------------------------------
    print("\n6. honest failures")
    missing = R.read(str(tmp / "there-is-no-such-file.pdf"))
    check("missing file: not ok", not missing.ok)
    check("missing file: says why", "no such file" in missing.warning_text().lower())
    check("missing file: never invents content", missing.text.strip() == "")
    summary = R.summarise(missing)
    check("missing file: summary refuses", "will not pretend" in summary)

    if "pdf_blank" in made:
        blank = R.read(str(made["pdf_blank"]))
        check("scanned pdf: warns about OCR",
              any("no text layer" in w for w in blank.warnings), blank.warning_text()[:60])
    else:
        skip("blank pdf", "pypdf not installed")

    empty = tmp / "empty.txt"
    empty.write_text("", encoding="utf-8")
    check("empty text file: summary says so",
          "almost no readable words" in R.summarise(R.read(str(empty))).lower())

    # 7. size guard -------------------------------------------------------
    print("\n7. limits")
    check("read limit is capped", R.MAX_READ <= 200_000, str(R.MAX_READ))
    check("byte limit is 64 MB", R.MAX_BYTES == 64 * 1024 * 1024, str(R.MAX_BYTES))

    print(f"\n{'ALL PASS' if not FAIL else 'FAILURES'} — {PASS} passed, {FAIL} failed, {SKIP} skipped")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
